import argparse
import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[1]

TEMPLATE_EXACTA = BASE_DIR / "manuals" / "04. FORMATO INSPECCION NIVEL 2 CON CALCULO DE IC SIN BLOQUEO (1).xlsm"
OUTPUT_DIR = BASE_DIR / "outputs"

COLUMNAS_SEVERIDAD = {
    "0": "F",
    "1": "G",
    "2": "H",
    "3": "I",
    "4": "J",
    "5": "K",
}

COLUMNA_FOTOS = "L"
COLUMNA_UBICACION = "M"


def encontrar_plantilla():
    if TEMPLATE_EXACTA.exists():
        return TEMPLATE_EXACTA

    candidatos = list((BASE_DIR / "manuals").glob("*.xlsm"))

    for archivo in candidatos:
        if "FORMATO INSPECCION" in archivo.name.upper():
            return archivo

    raise FileNotFoundError("No encontré la plantilla .xlsm en manuals.")


def texto(valor):
    return str(valor or "").strip()


def nombre_salida_seguro(data):
    nombre = texto(data.get("nombreArchivo"))

    if not nombre:
        fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
        nombre = f"inspeccion_sipucol_{fecha}.xlsm"

    if not nombre.lower().endswith(".xlsm"):
        nombre += ".xlsm"

    return nombre


def celda_real(ws, celda):
    for rango in ws.merged_cells.ranges:
        if celda.coordinate in rango:
            return ws.cell(rango.min_row, rango.min_col)

    return celda


def escribir_celda(ws, fila, columna, valor):
    celda = ws[f"{columna}{fila}"]
    celda_destino = celda_real(ws, celda)
    celda_destino.value = valor


def buscar_texto_y_escribir_derecha(ws, label, valor, ocurrencia=1):
    if valor in [None, ""]:
        return False

    label = label.lower()
    contador = 0

    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40)):
        for cell in row:
            contenido = texto(cell.value).lower()

            if label in contenido:
                contador += 1

                if contador != ocurrencia:
                    continue

                destino = ws.cell(cell.row, cell.column + 1)
                destino = celda_real(ws, destino)
                destino.value = valor
                return True

    return False


def llenar_identificacion(ws, identificacion):
    buscar_texto_y_escribir_derecha(ws, "Nombre del puente", identificacion.get("nombrePuente"))
    buscar_texto_y_escribir_derecha(ws, "ID Puente", identificacion.get("idPuente"))
    buscar_texto_y_escribir_derecha(ws, "Administrador vial", identificacion.get("administradorVial"))
    buscar_texto_y_escribir_derecha(ws, "Entidad administradora", identificacion.get("entidadAdministradora"))
    buscar_texto_y_escribir_derecha(ws, "Responsable del diligenciamiento", identificacion.get("responsableDiligenciamiento"))
    buscar_texto_y_escribir_derecha(ws, "Responsable de la revisión", identificacion.get("responsableRevision"))
    buscar_texto_y_escribir_derecha(ws, "Fecha de levantamiento", identificacion.get("fecha"))
    buscar_texto_y_escribir_derecha(ws, "Hora", identificacion.get("hora"))

    buscar_texto_y_escribir_derecha(ws, "Cargo", identificacion.get("cargoDiligenciamiento"), ocurrencia=1)
    buscar_texto_y_escribir_derecha(ws, "Cargo", identificacion.get("cargoRevision"), ocurrencia=2)

    buscar_texto_y_escribir_derecha(ws, "Tarjeta Profesional", identificacion.get("tarjetaDiligenciamiento"), ocurrencia=1)
    buscar_texto_y_escribir_derecha(ws, "Tarjeta Profesional", identificacion.get("tarjetaRevision"), ocurrencia=2)


def obtener_hoja(wb, nombre):
    nombre = texto(nombre).lower()

    for ws in wb.worksheets:
        if ws.title.strip().lower() == nombre:
            return ws

    return None


def fila_contiene_codigo(ws, fila, codigo):
    codigo = codigo.upper()

    for col in range(1, 8):
        valor = texto(ws.cell(fila, col).value).upper()

        if valor == codigo:
            return True

    return False


def buscar_fila_codigo(ws, codigo, fila_sugerida=None):
    codigo = texto(codigo).upper()

    if fila_sugerida:
        try:
            fila = int(fila_sugerida)

            if fila_contiene_codigo(ws, fila, codigo):
                return fila
        except Exception:
            pass

    for fila in range(1, ws.max_row + 1):
        if fila_contiene_codigo(ws, fila, codigo):
            return fila

    return None


def limpiar_severidad(ws, fila):
    for columna in COLUMNAS_SEVERIDAD.values():
        escribir_celda(ws, fila, columna, "")


def escribir_item(ws, item):
    codigo = texto(item.get("codigo"))

    if not codigo:
        return False

    fila = buscar_fila_codigo(ws, codigo, item.get("filaExcel"))

    if not fila:
        print(f"No encontré {codigo} en hoja {ws.title}")
        return False

    severidad = texto(item.get("severidad"))

    if severidad in COLUMNAS_SEVERIDAD:
        limpiar_severidad(ws, fila)
        escribir_celda(ws, fila, COLUMNAS_SEVERIDAD[severidad], "X")

    fotos = texto(item.get("fotos"))
    ubicacion = texto(item.get("ubicacion"))

    if fotos:
        escribir_celda(ws, fila, COLUMNA_FOTOS, fotos)

    if ubicacion:
        escribir_celda(ws, fila, COLUMNA_UBICACION, ubicacion)

    return True


def exportar_excel(data):
    plantilla = encontrar_plantilla()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    salida = OUTPUT_DIR / nombre_salida_seguro(data)

    shutil.copy2(plantilla, salida)

    wb = load_workbook(salida, keep_vba=True)

    identificacion = data.get("identificacion", {})

    for ws in wb.worksheets:
        llenar_identificacion(ws, identificacion)

    escritos = 0

    for componente in data.get("componentes", []):
        nombre_hoja = componente.get("nombre") or componente.get("hoja")
        ws = obtener_hoja(wb, nombre_hoja)

        if not ws:
            print(f"No encontré hoja: {nombre_hoja}")
            continue

        for item in componente.get("items", []):
            if escribir_item(ws, item):
                escritos += 1

    wb.save(salida)

    print("EXPORTACIÓN LISTA")
    print(f"Archivo creado: {salida}")
    print(f"Filas escritas: {escritos}")


def crear_data_prueba():
    return {
        "nombreArchivo": "PRUEBA_EXPORTACION_SIPUCOL.xlsm",
        "identificacion": {
            "nombrePuente": "PUENTE DE PRUEBA",
            "idPuente": "TEST-001",
            "administradorVial": "INVIAS",
            "entidadAdministradora": "ENTIDAD TEST",
            "responsableDiligenciamiento": "Inspector de prueba",
            "cargoDiligenciamiento": "Inspector",
            "tarjetaDiligenciamiento": "TP-123",
            "responsableRevision": "Revisor de prueba",
            "cargoRevision": "Revisor",
            "tarjetaRevision": "TP-456",
            "fecha": "16/06/2026",
            "hora": "10:30"
        },
        "componentes": [
            {
                "nombre": "Superficie de accesos",
                "items": [
                    {
                        "codigo": "D1DS2",
                        "filaExcel": "",
                        "severidad": "3",
                        "fotos": "1,2",
                        "ubicacion": "PRUEBA: daño localizado en acceso derecho."
                    }
                ]
            }
        ]
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--test", action="store_true")
    parser.add_argument("--json", type=str)

    args = parser.parse_args()

    if args.test:
        data = crear_data_prueba()
    elif args.json:
        with open(args.json, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        raise SystemExit("Usa --test o --json archivo.json")

    exportar_excel(data)


if __name__ == "__main__":
    main()
